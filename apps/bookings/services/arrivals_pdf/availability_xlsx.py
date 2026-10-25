"""Excel twin of the availability flyer — port blocks side by side, continuous dates."""

from __future__ import annotations

from datetime import time
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.bookings.services.arrivals_pdf.availability_block import (
    AvailabilityPortBlock,
    build_port_blocks,
)
from apps.bookings.services.arrivals_pdf.labels import (
    AVAILABILITY_HEADING,
    FOOTER_NOTE_CAL,
    range_span_label,
    updated_label,
)
from apps.bookings.services.arrivals_pdf.logo import add_worksheet_port_logo
from apps.bookings.services.arrivals_pdf.overview import ExportOverview, paint_overview
from apps.bookings.services.arrivals_pdf.occupancy import (
    occupancy_cell_text,
    occupancy_date_label,
    occupancy_weekday,
)

NAVY = "1F4E79"
INK = "111111"
WHITE = "FFFFFF"
GRID = "1A1A1A"
MUTED = "4B5563"
ALT = "F8FAFC"

FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_WHITE = PatternFill("solid", fgColor=WHITE)
FILL_ALT = PatternFill("solid", fgColor=ALT)
GAP_COLS = 1
MIN_BLOCK_WIDTH = 74
FECHA_W = 14
DIA_W = 10
POS_W = 18
FONT_TITLE = Font(name="Cambria", size=13, bold=True, color=INK)
FONT_RANGE = Font(name="Cambria", size=11, bold=True, color=INK)
FONT_SUB = Font(name="Calibri", size=9, italic=True, color=MUTED)
FONT_HEAD = Font(name="Calibri", size=9, bold=True, color=WHITE)
FONT_BODY = Font(name="Calibri", size=8, color=INK)
FONT_FOOTER = Font(name="Calibri", size=8, color=WHITE)
ALIGN_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(style="thin", color=GRID)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_arrivals_availability_xlsx(
    bookings: list,
    *,
    ports: list,
    date_from,
    date_to,
    overview: ExportOverview | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Disponibilidad"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    updated = timezone.localdate()
    blocks = build_port_blocks(bookings, ports, date_from, date_to)
    span = 8
    if blocks:
        widths = [block.col_count for block in blocks]
        span = sum(widths) + GAP_COLS * max(len(widths) - 1, 0)
    band_row = paint_overview(ws, overview, last_col=max(span, 8))
    start_col = 1
    for block in blocks:
        _paint_port(
            ws,
            block=block,
            date_from=date_from,
            date_to=date_to,
            updated=updated,
            start_row=band_row,
            start_col=start_col,
        )
        start_col += block.col_count + GAP_COLS
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _paint_port(
    ws,
    *,
    block: AvailabilityPortBlock,
    date_from,
    date_to,
    updated,
    start_row: int,
    start_col: int,
) -> int:
    labels = block.labels or ["—"]
    n_pos = max(len(labels), 1)
    last_col = start_col + 1 + n_pos
    port = block.port
    port_name = (getattr(port, "name", None) or "Puerto").strip()
    pos_w = POS_W
    total_w = FECHA_W + DIA_W + pos_w * n_pos
    if total_w < MIN_BLOCK_WIDTH:
        pos_w += (MIN_BLOCK_WIDTH - total_w) / n_pos
    ws.column_dimensions[get_column_letter(start_col)].width = FECHA_W
    ws.column_dimensions[get_column_letter(start_col + 1)].width = DIA_W
    for i in range(n_pos):
        ws.column_dimensions[get_column_letter(start_col + 2 + i)].width = pos_w

    heading_row = start_row
    range_row = start_row + 1
    sub_row = start_row + 2
    has_logo = add_worksheet_port_logo(
        ws, port, row=heading_row, col=start_col, max_w=72, max_h=32
    )
    title_start = start_col + 1 if has_logo else start_col
    ws.merge_cells(
        start_row=heading_row,
        start_column=title_start,
        end_row=heading_row,
        end_column=last_col,
    )
    heading = ws.cell(heading_row, title_start, AVAILABILITY_HEADING)
    heading.font = FONT_TITLE
    heading.alignment = ALIGN_C
    ws.row_dimensions[heading_row].height = 22
    ws.merge_cells(
        start_row=range_row,
        start_column=title_start,
        end_row=range_row,
        end_column=last_col,
    )
    span = ws.cell(range_row, title_start, range_span_label(date_from, date_to))
    span.font = FONT_RANGE
    span.alignment = ALIGN_C
    ws.row_dimensions[range_row].height = 18
    ws.merge_cells(
        start_row=sub_row, start_column=start_col, end_row=sub_row, end_column=last_col
    )
    c2 = ws.cell(
        sub_row,
        start_col,
        f"{port_name}  ·  La asignación de posición está sujeta a cambios",
    )
    c2.font = FONT_SUB
    c2.alignment = ALIGN_C
    ws.row_dimensions[sub_row].height = 16

    head_row = start_row + 3
    headers = ["Fecha", "Día", *labels]
    for offset, label in enumerate(headers):
        cell = ws.cell(head_row, start_col + offset, label)
        cell.font = FONT_HEAD
        cell.fill = FILL_NAVY
        cell.alignment = ALIGN_C
        cell.border = BORDER
    count_row = start_row + 4
    padded = list(block.counts) + [0] * max(0, n_pos - len(block.counts))
    count_vals = ["Arribos", "", *[str(n) for n in padded[:n_pos]]]
    for offset, value in enumerate(count_vals):
        cell = ws.cell(count_row, start_col + offset, value)
        cell.font = FONT_HEAD
        cell.fill = FILL_NAVY
        cell.alignment = ALIGN_C
        cell.border = BORDER

    body_start = start_row + 5
    for index, day in enumerate(block.days):
        r = body_start + index
        fill = FILL_ALT if index % 2 else FILL_WHITE
        ws.row_dimensions[r].height = 28
        cells = [""] * n_pos
        for booking in sorted(
            block.occupied.get(day, []),
            key=lambda b: (b.eta or time.min, (b.vessel.name if b.vessel_id else "")),
        ):
            text = occupancy_cell_text(booking)
            idx = block.pos_index.get(booking.position_id)
            if idx is None:
                idx = n_pos - 1
            cells[idx] = f"{cells[idx]}\n{text}".strip() if cells[idx] else text
        values = [occupancy_date_label(day), occupancy_weekday(day), *cells]
        for offset, value in enumerate(values):
            cell = ws.cell(r, start_col + offset, value)
            cell.font = FONT_BODY
            cell.fill = fill
            cell.alignment = ALIGN_C if offset > 0 else ALIGN_L
            cell.border = BORDER

    footer_row = body_start + len(block.days)
    ws.merge_cells(
        start_row=footer_row,
        start_column=start_col,
        end_row=footer_row,
        end_column=last_col,
    )
    footer = ws.cell(
        footer_row,
        start_col,
        f"{port_name} · {FOOTER_NOTE_CAL} {updated_label(updated)}",
    )
    footer.font = FONT_FOOTER
    footer.fill = FILL_NAVY
    footer.alignment = ALIGN_C
    footer.border = BORDER
    for col in range(start_col + 1, last_col + 1):
        extra = ws.cell(footer_row, col)
        extra.fill = FILL_NAVY
        extra.border = BORDER
    ws.row_dimensions[footer_row].height = 18
    return footer_row
