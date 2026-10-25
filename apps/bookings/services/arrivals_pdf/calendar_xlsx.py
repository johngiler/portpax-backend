"""Excel twin of the Beto occupancy calendar — months side by side."""

from __future__ import annotations

from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.bookings.services.arrivals_pdf.calendar_block import (
    BLOCK_COLS,
    MonthBlock,
    WeekBand,
    build_month_block,
    build_week_block,
    week_dates,
)
from apps.bookings.services.arrivals_pdf.labels import (
    FOOTER_NOTE_CAL,
    arrivals_range_title,
    updated_label,
)
from apps.bookings.services.arrivals_pdf.logo import add_worksheet_port_logo
from apps.bookings.services.arrivals_pdf.list_data import group_bookings_by_port
from apps.bookings.services.arrivals_pdf.month_grid import (
    excel_sheet_title,
    group_port_months,
    months_in_range,
)
from apps.bookings.services.arrivals_pdf.overview import ExportOverview, paint_overview

MONTHS_PER_ROW = 3
GAP_COLS = 1

NAVY = "1F4E79"
STEEL = "5B9BD5"
INK = "111111"
WHITE = "FFFFFF"
GRID = "1A1A1A"

FILL_STEEL = PatternFill("solid", fgColor=STEEL)
FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_WHITE = PatternFill("solid", fgColor=WHITE)
FILL_DATE = PatternFill("solid", fgColor="E8F0F7")
FONT_TITLE = Font(name="Cambria", size=16, bold=True, color=INK)
FONT_MONTH = Font(name="Calibri", size=10, bold=True, color=WHITE)
FONT_WEEK = Font(name="Calibri", size=8, bold=True, color=WHITE)
FONT_DATE = Font(name="Calibri", size=7, color=INK)
FONT_POS = Font(name="Calibri", size=8, bold=True, color=INK)
FONT_SHIP = Font(name="Calibri", size=7, color=INK)
FONT_FOOTER = Font(name="Calibri", size=8, color=WHITE)
ALIGN_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color=GRID)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_arrivals_calendar_xlsx(
    bookings: list,
    *,
    overview: ExportOverview | None = None,
    mode: str = "monthly",
    date_from=None,
    date_to=None,
    ports: list | None = None,
) -> bytes:
    if mode == "weekly" and date_from is not None:
        return _build_week_xlsx(
            bookings,
            overview=overview,
            date_from=date_from,
            date_to=date_to,
            ports=ports,
        )
    if mode != "annual":
        return _build_month_xlsx(
            bookings,
            overview=overview,
            date_from=date_from,
            date_to=date_to,
            ports=ports,
        )
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    updated = timezone.localdate()
    by_port: dict[int, dict] = {}
    order: list[int] = []
    for port, year, month, rows in group_port_months(bookings):
        if port.id not in by_port:
            by_port[port.id] = {"port": port, "months": []}
            order.append(port.id)
        by_port[port.id]["months"].append((year, month, rows))
    for port_id in order:
        pack = by_port[port_id]
        port = pack["port"]
        ws = wb.create_sheet(_unique_sheet(used, port.name or "Puerto"))
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        dates = []
        for year, month, rows in pack["months"]:
            dates.extend(b.call_date for b in rows)
        span = MONTHS_PER_ROW * (BLOCK_COLS + GAP_COLS) - GAP_COLS
        start_row = paint_overview(ws, overview, last_col=span)
        _paint_sheet_banner(
            ws,
            port=port,
            title=arrivals_range_title(min(dates), max(dates)),
            span=span,
            start_row=start_row,
        )
        col_i = 0
        band_row = start_row + 2
        band_bottom = band_row
        for year, month, rows in pack["months"]:
            if col_i == MONTHS_PER_ROW:
                band_row = band_bottom + 2
                col_i = 0
            start_col = 1 + col_i * (BLOCK_COLS + GAP_COLS)
            block = build_month_block(port, year, month, rows)
            bottom = _paint_month(
                ws, block=block, updated=updated, start_row=band_row, start_col=start_col
            )
            band_bottom = max(band_bottom, bottom)
            col_i += 1
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _build_week_xlsx(
    bookings: list,
    *,
    overview: ExportOverview | None,
    date_from,
    date_to,
    ports: list | None,
) -> bytes:
    days = week_dates(date_from, date_to)
    grouped = group_bookings_by_port(bookings)
    by_port = {port.id: rows for port, rows in grouped}
    port_list = list(ports or [port for port, _rows in grouped])
    items = [
        (port, build_week_block(port, days, by_port.get(port.id, [])))
        for port in port_list
    ]
    return _workbook_one_sheet(
        items, overview=overview, sheet_title="Semana", row_h=22
    )


def _build_month_xlsx(
    bookings: list,
    *,
    overview: ExportOverview | None,
    date_from,
    date_to,
    ports: list | None,
) -> bytes:
    grouped = group_bookings_by_port(bookings)
    by_port = {port.id: rows for port, rows in grouped}
    port_list = list(ports or [port for port, _rows in grouped])
    if date_from is not None and date_to is not None:
        months = months_in_range(date_from, date_to)
    else:
        months = []
        seen: set[tuple[int, int]] = set()
        for booking in bookings:
            key = (booking.call_date.year, booking.call_date.month)
            if key not in seen:
                seen.add(key)
                months.append(key)
        months.sort()
    items: list[tuple[object, MonthBlock]] = []
    for port in port_list:
        rows = by_port.get(port.id, [])
        for year, month in months:
            month_rows = [
                booking
                for booking in rows
                if booking.call_date.year == year and booking.call_date.month == month
            ]
            items.append((port, build_month_block(port, year, month, month_rows)))
    return _workbook_one_sheet(items, overview=overview, sheet_title="Mes", row_h=15)


def _workbook_one_sheet(
    items: list[tuple[object, MonthBlock]],
    *,
    overview: ExportOverview | None,
    sheet_title: str,
    row_h: int,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = excel_sheet_title(sheet_title)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    updated = timezone.localdate()
    span = min(len(items), MONTHS_PER_ROW) * (BLOCK_COLS + GAP_COLS) - GAP_COLS
    start_row = paint_overview(ws, overview, last_col=max(span, BLOCK_COLS))
    col_i = 0
    band_row = start_row
    band_bottom = start_row
    start_col = 1
    for port, block in items:
        if col_i == MONTHS_PER_ROW:
            band_row = band_bottom + 2
            col_i = 0
            start_col = 1
        has_logo = add_worksheet_port_logo(
            ws, port, row=band_row, col=start_col, max_w=88, max_h=32
        )
        last_col = start_col + BLOCK_COLS - 1
        title_col = start_col + 2 if has_logo else start_col
        ws.merge_cells(
            start_row=band_row,
            start_column=title_col,
            end_row=band_row,
            end_column=last_col,
        )
        cell = ws.cell(
            band_row,
            title_col,
            f"{(getattr(port, 'name', None) or '').strip()}  ·  {block.title}",
        )
        cell.font = FONT_TITLE
        cell.alignment = ALIGN_C
        ws.row_dimensions[band_row].height = 36 if has_logo else 22
        bottom = _paint_month(
            ws,
            block=block,
            updated=updated,
            start_row=band_row + 1,
            start_col=start_col,
            row_h=row_h,
        )
        band_bottom = max(band_bottom, bottom)
        start_col += BLOCK_COLS + GAP_COLS
        col_i += 1
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _unique_sheet(used: set[str], raw: str) -> str:
    base = excel_sheet_title(raw)
    name = base
    n = 2
    while name in used:
        suffix = f" {n}"
        name = excel_sheet_title(base[: 31 - len(suffix)] + suffix)
        n += 1
    used.add(name)
    return name


def _paint_sheet_banner(ws, *, port, title: str, span: int, start_row: int = 1) -> None:
    last = max(span, BLOCK_COLS)
    has_logo = add_worksheet_port_logo(
        ws, port, row=start_row, col=1, max_w=100, max_h=40
    )
    title_col = 3 if has_logo else 1
    ws.merge_cells(
        start_row=start_row, start_column=title_col, end_row=start_row, end_column=last
    )
    cell = ws.cell(start_row, title_col, f"{(port.name or '').strip()}  ·  {title}")
    cell.font = FONT_TITLE
    cell.alignment = ALIGN_C
    ws.row_dimensions[start_row].height = 40 if has_logo else 24
    ws.row_dimensions[start_row + 1].height = 8


def _paint_month(
    ws,
    *,
    block: MonthBlock,
    updated,
    start_row: int,
    start_col: int,
    row_h: int = 15,
) -> int:
    last_col = start_col + BLOCK_COLS - 1
    widths = (5, 8, 13, 13, 13, 13, 13, 13, 13)
    for offset, width in enumerate(widths):
        ws.column_dimensions[get_column_letter(start_col + offset)].width = width

    head = start_row
    _fill_cell(ws, head, start_col, str(block.total), FONT_WEEK, FILL_STEEL, ALIGN_C)
    _fill_cell(ws, head, start_col + 1, "Pos.", FONT_WEEK, FILL_STEEL, ALIGN_C)
    for i, label in enumerate(block.weekday_labels):
        _fill_cell(ws, head, start_col + 2 + i, label, FONT_WEEK, FILL_STEEL, ALIGN_C)
    ws.row_dimensions[head].height = 16

    title_row = start_row + 1
    ws.merge_cells(
        start_row=title_row,
        start_column=start_col,
        end_row=title_row,
        end_column=last_col,
    )
    _fill_span(
        ws, title_row, start_col, last_col, FILL_STEEL, FONT_MONTH, ALIGN_C, block.title
    )
    ws.row_dimensions[title_row].height = 16

    row = title_row + 1
    for week in block.weeks:
        row = _paint_week(ws, row, start_col, week, row_h=row_h)
    ws.merge_cells(
        start_row=row, start_column=start_col, end_row=row, end_column=last_col
    )
    _fill_span(
        ws,
        row,
        start_col,
        last_col,
        FILL_NAVY,
        FONT_FOOTER,
        ALIGN_C,
        f"{FOOTER_NOTE_CAL} {updated_label(updated)}",
    )
    ws.row_dimensions[row].height = 16
    return row


def _paint_week(ws, row: int, start_col: int, week: WeekBand, row_h: int = 15) -> int:
    _fill_cell(ws, row, start_col, str(week.count), FONT_POS, FILL_DATE, ALIGN_C)
    _fill_cell(ws, row, start_col + 1, "", FONT_DATE, FILL_DATE, ALIGN_C)
    for i, stamp in enumerate(week.dates):
        _fill_cell(ws, row, start_col + 2 + i, stamp, FONT_DATE, FILL_DATE, ALIGN_C)
    ws.row_dimensions[row].height = 14
    row += 1
    for label, cells in week.rows:
        _fill_cell(ws, row, start_col, "", FONT_SHIP, FILL_WHITE, ALIGN_C)
        _fill_cell(ws, row, start_col + 1, label, FONT_POS, FILL_WHITE, ALIGN_C)
        for i, value in enumerate(cells):
            _fill_cell(ws, row, start_col + 2 + i, value, FONT_SHIP, FILL_WHITE, ALIGN_C)
        ws.row_dimensions[row].height = row_h
        row += 1
    return row


def _fill_cell(ws, row, col, value, font, fill, alignment) -> None:
    cell = ws.cell(row, col, value)
    cell.font = font
    cell.fill = fill
    cell.alignment = alignment
    cell.border = BORDER


def _fill_span(ws, row, start, end, fill, font, alignment, value="") -> None:
    for col in range(start, end + 1):
        cell = ws.cell(row, col, value if col == start else None)
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment
        cell.border = BORDER
