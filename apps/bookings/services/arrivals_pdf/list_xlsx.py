"""Excel twin of the list PDF — same columns, port blocks side by side."""

from __future__ import annotations

from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.bookings.services.arrivals_pdf.labels import arrivals_range_title
from apps.bookings.services.arrivals_pdf.list_data import (
    group_bookings_by_port,
    list_center_cols,
    list_col_widths,
    list_headers,
    list_row_values,
)
from apps.bookings.services.arrivals_pdf.logo import add_worksheet_port_logo
from apps.bookings.services.arrivals_pdf.overview import ExportOverview, paint_overview
from apps.bookings.services.report_exports.common import PAX_BASIS_PLANNED

NAVY = "1F4E79"
INK = "111111"
WHITE = "FFFFFF"
GRID = "1A1A1A"

FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_WHITE = PatternFill("solid", fgColor=WHITE)
FONT_TITLE = Font(name="Cambria", size=14, bold=True, color=INK)
FONT_HEAD = Font(name="Calibri", size=9, bold=True, color=WHITE)
FONT_BODY = Font(name="Calibri", size=8, color=INK)
ALIGN_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(style="thin", color=GRID)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

GAP_COLS = 1


def build_arrivals_list_xlsx(
    bookings: list,
    *,
    overview: ExportOverview | None = None,
    pax_basis: str = PAX_BASIS_PLANNED,
    include_cancellation: bool = False,
) -> bytes:
    headers = list_headers(include_cancellation=include_cancellation)
    block_cols = len(headers)
    wb = Workbook()
    ws = wb.active
    ws.title = "Arribos"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    groups = group_bookings_by_port(bookings)
    n_blocks = max(len(groups), 1)
    last_col = n_blocks * (block_cols + GAP_COLS) - GAP_COLS
    start_row = paint_overview(ws, overview, last_col=last_col)
    start_col = 1
    for port, rows in groups:
        _paint_port_block(
            ws,
            start_col=start_col,
            start_row=start_row,
            port=port,
            rows=rows,
            pax_basis=pax_basis,
            include_cancellation=include_cancellation,
        )
        start_col += block_cols + GAP_COLS
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _paint_port_block(
    ws,
    *,
    start_col: int,
    start_row: int = 1,
    port,
    rows: list,
    pax_basis: str = PAX_BASIS_PLANNED,
    include_cancellation: bool = False,
) -> None:
    headers = list_headers(include_cancellation=include_cancellation)
    widths = list_col_widths(include_cancellation=include_cancellation)
    center = list_center_cols(include_cancellation=include_cancellation)
    last_col = start_col + len(headers) - 1
    for offset, width in enumerate(widths):
        ws.column_dimensions[get_column_letter(start_col + offset)].width = width

    has_logo = add_worksheet_port_logo(
        ws, port, row=start_row, col=start_col, max_w=88, max_h=36
    )
    title_col = start_col + 2 if has_logo else start_col
    start = rows[0].call_date if rows else timezone.localdate()
    end = rows[-1].call_date if rows else start
    title = f"{arrivals_range_title(start, end)}  ·  {(port.name or '').strip()}"
    ws.merge_cells(
        start_row=start_row, start_column=title_col, end_row=start_row, end_column=last_col
    )
    cell = ws.cell(start_row, title_col, title)
    cell.font = FONT_TITLE
    cell.alignment = ALIGN_C
    ws.row_dimensions[start_row].height = 40 if has_logo else 22

    head_row = start_row + 1
    for offset, header in enumerate(headers):
        col = start_col + offset
        head = ws.cell(head_row, col, header)
        head.font = FONT_HEAD
        head.fill = FILL_NAVY
        head.alignment = ALIGN_C
        head.border = BORDER
    ws.row_dimensions[head_row].height = 18

    for index, booking in enumerate(rows, start=1):
        excel_row = head_row + index
        ws.row_dimensions[excel_row].height = 18 if include_cancellation else 16
        values = list_row_values(
            booking,
            number=index,
            pax_basis=pax_basis,
            include_cancellation=include_cancellation,
        )
        for offset, value in enumerate(values):
            col = start_col + offset
            body = ws.cell(excel_row, col, value)
            body.font = FONT_BODY
            body.fill = FILL_WHITE
            body.alignment = ALIGN_C if offset in center else ALIGN_L
            body.border = BORDER
