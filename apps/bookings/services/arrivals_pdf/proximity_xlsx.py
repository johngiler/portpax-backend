"""Excel twin of the proximity flyer — dates × ports, plus a detail sheet."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.bookings.services.arrivals_pdf.labels import format_hhmm, list_date_label
from apps.bookings.services.arrivals_pdf.logo import add_worksheet_port_logo, load_ports_by_ids
from apps.bookings.services.arrivals_pdf.overview import ExportOverview, paint_overview
from apps.bookings.services.booking_export import STATUS_LABELS_ES

NAVY = "1F4E79"
INK = "111111"
WHITE = "FFFFFF"
GRID = "1A1A1A"
MUTED = "4B5563"
SAME = "DBEAFE"
PROX = "FEF3C7"
ALT = "F8FAFC"

FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_WHITE = PatternFill("solid", fgColor=WHITE)
FILL_SAME = PatternFill("solid", fgColor=SAME)
FILL_PROX = PatternFill("solid", fgColor=PROX)
FILL_ALT = PatternFill("solid", fgColor=ALT)
FONT_TITLE = Font(name="Calibri", size=9, bold=True, color=WHITE)
FONT_DATE = Font(name="Calibri", size=8, color=INK)
FONT_CELL = Font(name="Calibri", size=7, color=INK)
FONT_HEAD = Font(name="Calibri", size=9, bold=True, color=WHITE)
FONT_BODY = Font(name="Calibri", size=8, color=INK)
ALIGN_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(style="thin", color=GRID)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

PORTS_PER_ROW = 6
DATE_W = 16
PORT_W = 18
DETAIL_HEADERS = (
    "N/M",
    "Fecha",
    "Puerto",
    "Posición",
    "Código",
    "Estado",
    "ETA",
    "ETD",
    "Alerta",
    "1er arribo",
)


def build_proximity_matrix_xlsx(
    payload: dict, *, overview: ExportOverview | None = None
) -> bytes:
    ports = list(payload.get("ports") or []) or [{"id": 0, "name": "—"}]
    dates = list(payload.get("dates") or [])
    cells = list(payload.get("cells") or [])
    by_key = {
        (str(cell.get("date") or ""), int(cell.get("port_id") or 0)): cell
        for cell in cells
    }
    found_ports = {
        port.id: port
        for port in load_ports_by_ids([int(item.get("id") or 0) for item in ports])
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"

    bands = [
        ports[i : i + PORTS_PER_ROW] for i in range(0, len(ports), PORTS_PER_ROW)
    ] or [ports]
    span = max((len(band) + 1 for band in bands), default=2)
    start_row = paint_overview(ws, overview, last_col=max(span, 8))
    for band in bands:
        start_row = (
            _paint_matrix(
                ws,
                start_row=start_row,
                start_col=1,
                ports=band,
                dates=dates,
                by_key=by_key,
                found_ports=found_ports,
            )
            + 2
        )

    _paint_detail(wb.create_sheet("Escalas"), cells)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _paint_matrix(
    ws,
    *,
    start_row: int,
    start_col: int,
    ports: list,
    dates: list,
    by_key: dict,
    found_ports: dict,
) -> int:
    last_col = start_col + len(ports)
    ws.column_dimensions[get_column_letter(start_col)].width = DATE_W
    for offset, port in enumerate(ports):
        col = start_col + 1 + offset
        ws.column_dimensions[get_column_letter(col)].width = PORT_W
        add_worksheet_port_logo(
            ws,
            found_ports.get(int(port.get("id") or 0)),
            row=start_row,
            col=col,
            max_w=72,
            max_h=28,
        )
    ws.row_dimensions[start_row].height = 32
    head = start_row + 1
    _cell(ws, head, start_col, "Fecha", FONT_TITLE, FILL_NAVY, ALIGN_C)
    for offset, port in enumerate(ports):
        _cell(
            ws,
            head,
            start_col + 1 + offset,
            str(port.get("name") or ""),
            FONT_TITLE,
            FILL_NAVY,
            ALIGN_C,
        )
    ws.row_dimensions[head].height = 18
    body = head + 1
    for index, iso in enumerate(dates):
        row = body + index
        _cell(ws, row, start_col, _date_label(iso), FONT_DATE, FILL_WHITE, ALIGN_L)
        for offset, port in enumerate(ports):
            cell = by_key.get((iso, int(port.get("id") or 0)))
            fill = FILL_WHITE
            text = ""
            if cell:
                status = str(cell.get("cell_status") or "")
                if status == "same_day":
                    fill = FILL_SAME
                elif status == "proximity":
                    fill = FILL_PROX
                else:
                    fill = FILL_ALT
                text = _cell_text(cell)
            _cell(ws, row, start_col + 1 + offset, text, FONT_CELL, fill, ALIGN_C)
        ws.row_dimensions[row].height = 36
    return body + len(dates) - 1 if dates else head


def _paint_detail(ws, cells: list) -> None:
    ws.sheet_view.showGridLines = False
    widths = (6, 16, 16, 10, 16, 12, 8, 8, 14, 12)
    for offset, header in enumerate(DETAIL_HEADERS):
        ws.column_dimensions[get_column_letter(offset + 1)].width = widths[offset]
        _cell(ws, 1, offset + 1, header, FONT_HEAD, FILL_NAVY, ALIGN_C)
    ws.row_dimensions[1].height = 18
    ordered = sorted(
        cells,
        key=lambda cell: (str(cell.get("date") or ""), str(cell.get("port_name") or "")),
    )
    for index, cell in enumerate(ordered, start=1):
        values = (
            str(index),
            _date_label(str(cell.get("date") or "")),
            str(cell.get("port_name") or ""),
            str(cell.get("position_code") or "—"),
            str(cell.get("booking_code") or ""),
            STATUS_LABELS_ES.get(str(cell.get("status") or ""), str(cell.get("status") or "")),
            _iso_time(cell.get("eta")),
            _iso_time(cell.get("etd")),
            _alert_label(cell),
            "Sí" if cell.get("first_arrival") else "",
        )
        for offset, value in enumerate(values):
            _cell(ws, index + 1, offset + 1, value, FONT_BODY, FILL_WHITE, ALIGN_C if offset != 2 else ALIGN_L)
        ws.row_dimensions[index + 1].height = 16


def _cell(ws, row, col, value, font, fill, alignment) -> None:
    cell = ws.cell(row, col, value)
    cell.font = font
    cell.fill = fill
    cell.alignment = alignment
    cell.border = BORDER


def _cell_text(cell: dict) -> str:
    position = str(cell.get("position_code") or "").strip() or "—"
    times = "–".join(part for part in (_iso_time(cell.get("eta")), _iso_time(cell.get("etd"))) if part)
    code = str(cell.get("booking_code") or "").strip()
    alert = _alert_label(cell)
    lines = [position]
    if times:
        lines.append(times)
    if code:
        lines.append(code)
    if alert:
        lines.append(alert)
    return "\n".join(lines)


def _alert_label(cell: dict) -> str:
    status = str(cell.get("cell_status") or "")
    if status == "same_day":
        return "Mismo día"
    if status == "proximity":
        return "Proximidad"
    return ""


def _date_label(iso: str) -> str:
    try:
        return list_date_label(date.fromisoformat(iso))
    except ValueError:
        return iso


def _iso_time(value) -> str:
    if not value:
        return ""
    text = str(value)
    if len(text) >= 5 and text[2] == ":":
        return text[:5]
    try:
        from datetime import time as time_cls

        return format_hhmm(time_cls.fromisoformat(text))
    except ValueError:
        return text
