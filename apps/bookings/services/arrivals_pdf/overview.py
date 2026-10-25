"""Shared export overview: totals + filter criteria for bookings flyers."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib.colors import HexColor
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas

from apps.bookings.services.arrivals_pdf.chrome import (
    MUTED,
    NAVY_DEEP,
    STEEL,
    fit_text,
)
from apps.bookings.services.arrivals_pdf.labels import range_span_label
from apps.bookings.services.booking_export import STATUS_LABELS_ES
from apps.bookings.services.report_exports.common import (
    PAX_BASIS_CAPACITY,
    PAX_BASIS_PLANNED,
    booking_pax,
    format_pax_count,
    normalize_pax_basis,
)

KIND_LABELS = {
    "list": "Lista de arribos",
    "calendar": "Calendario de ocupación",
    "availability": "Disponibilidad de posiciones",
    "proximity": "Proximidad barco / puerto",
}

STATUS_FILTER_LABELS = {
    **STATUS_LABELS_ES,
    "completed": "Completada",
    "action": "Por acción",
}
CONFLICT_LABELS = {
    "yellow": "Conflicto amarillo",
    "red": "Conflicto rojo",
    "green": "Conflicto verde",
    "proximity": "Tipo · Proximidad",
    "loa": "Tipo · Eslora",
    "schedule": "Tipo · Horario",
    "position": "Tipo · Posición",
    "lta": "Tipo · LTA",
    "physical": "Tipo · Físico",
}
FILL_OVERVIEW = PatternFill("solid", fgColor="EEF4FA")
FONT_KIND = Font(name="Calibri", size=12, bold=True, color="163556")
FONT_KPI = Font(name="Calibri", size=10, bold=True, color="163556")
FONT_CRITERIA = Font(name="Calibri", size=8, italic=True, color="4B5563")
ALIGN_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
OVERVIEW_BORDER = Border(
    left=Side(style="thin", color="5B9BD5"),
    right=Side(style="thin", color="5B9BD5"),
    top=Side(style="thin", color="5B9BD5"),
    bottom=Side(style="thin", color="5B9BD5"),
)
OVERVIEW_H = 0.95 * cm
OVERVIEW_ROWS = 3


@dataclass(frozen=True)
class ExportKpi:
    label: str
    value: str


@dataclass(frozen=True)
class ExportOverview:
    kind: str
    kind_label: str
    kpis: list[ExportKpi] = field(default_factory=list)
    criteria: list[tuple[str, str]] = field(default_factory=list)
    legend: str = ""

    @property
    def kpi_line(self) -> str:
        return "    ".join(f"{item.label}  {item.value}" for item in self.kpis)

    @property
    def criteria_line(self) -> str:
        if not self.criteria:
            return "Filtros: sin filtros adicionales"
        joined = "  ·  ".join(f"{label}: {value}" for label, value in self.criteria)
        return f"Filtros: {joined}"


def build_overview(
    *,
    kind: str,
    query=None,
    bookings: list | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    extra_kpis: list[ExportKpi] | None = None,
    extra_criteria: list[tuple[str, str]] | None = None,
    payload: dict | None = None,
    pax_basis: str = PAX_BASIS_PLANNED,
) -> ExportOverview:
    rows = list(bookings or [])
    raw_basis = pax_basis
    if query is not None:
        raw_basis = _get(query, "pax_basis") or pax_basis
    basis = normalize_pax_basis(raw_basis)
    start, end = _range(rows, date_from, date_to, query)
    criteria = _criteria_from_query(query, start=start, end=end, pax_basis=basis)
    if extra_criteria:
        criteria.extend(extra_criteria)
    kpis = extra_kpis or _default_kpis(kind, rows, payload, pax_basis=basis)
    return ExportOverview(
        kind=kind,
        kind_label=KIND_LABELS.get(kind, "Reservas"),
        kpis=kpis,
        criteria=_dedupe_criteria(criteria),
        legend=_legend_for(kind),
    )


def draw_overview_band(
    pdf: canvas.Canvas,
    x: float,
    y: float,
    width: float,
    overview: ExportOverview | None,
) -> float:
    if overview is None:
        return y
    h = OVERVIEW_H
    pdf.setFillColor(HexColor("#EEF4FA"))
    pdf.rect(x, y - h, width, h, fill=1, stroke=0)
    pdf.setStrokeColor(STEEL)
    pdf.setLineWidth(0.7)
    pdf.rect(x, y - h, width, h, fill=0, stroke=1)
    pdf.setFillColor(NAVY_DEEP)
    pdf.setFont("Helvetica-Bold", 8)
    pdf.drawString(x + 6, y - 12, overview.kind_label)
    if overview.kpi_line:
        pdf.setFont("Helvetica-Bold", 8.5)
        pdf.drawRightString(x + width - 6, y - 12, overview.kpi_line)
    pdf.setFillColor(MUTED)
    pdf.setFont("Helvetica", 6.6)
    line = overview.criteria_line
    if overview.legend:
        line = f"{line}    {overview.legend}"
    pdf.drawString(x + 6, y - 24, fit_text(pdf, line, width - 12))
    return y - h - 4


def paint_overview(
    ws,
    overview: ExportOverview | None,
    *,
    start_row: int = 1,
    start_col: int = 1,
    last_col: int = 8,
) -> int:
    if overview is None:
        return start_row
    last = max(last_col, start_col)
    for offset, value, font, height in (
        (0, overview.kind_label, FONT_KIND, 20),
        (1, overview.kpi_line, FONT_KPI, 18),
        (2, overview.criteria_line, FONT_CRITERIA, 28),
    ):
        row = start_row + offset
        ws.merge_cells(
            start_row=row, start_column=start_col, end_row=row, end_column=last
        )
        cell = ws.cell(row, start_col, value)
        cell.font = font
        cell.fill = FILL_OVERVIEW
        cell.alignment = ALIGN_L
        cell.border = OVERVIEW_BORDER
        for col in range(start_col + 1, last + 1):
            extra = ws.cell(row, col)
            extra.fill = FILL_OVERVIEW
            extra.border = OVERVIEW_BORDER
        ws.row_dimensions[row].height = height
    return start_row + OVERVIEW_ROWS + 1


def _legend_for(kind: str) -> str:
    if kind == "proximity":
        return "Leyenda: azul = mismo día · ámbar = proximidad"
    return ""


def _range(rows, date_from, date_to, query) -> tuple[date | None, date | None]:
    start = date_from or _query_date(query, "date_from", "call_date_from")
    end = date_to or _query_date(query, "date_to", "call_date_to")
    dates = [item.call_date for item in rows if getattr(item, "call_date", None)]
    if dates:
        start = start or min(dates)
        end = end or max(dates)
    return start, end


def _default_kpis(
    kind: str,
    rows: list,
    payload: dict | None,
    *,
    pax_basis: str = PAX_BASIS_PLANNED,
) -> list[ExportKpi]:
    if kind == "proximity" and payload is not None:
        cells = list(payload.get("cells") or [])
        ports = list(payload.get("ports") or [])
        dates = list(payload.get("dates") or [])
        same = sum(1 for cell in cells if cell.get("cell_status") == "same_day")
        prox = sum(1 for cell in cells if cell.get("cell_status") == "proximity")
        return [
            ExportKpi("Escalas", str(len(cells))),
            ExportKpi("Puertos", str(len(ports))),
            ExportKpi("Días", str(len(dates))),
            ExportKpi("Mismo día", str(same)),
            ExportKpi("Proximidad", str(prox)),
        ]
    ports = {getattr(item.port, "name", "") for item in rows if getattr(item, "port_id", None)}
    lines = {
        getattr(item.shipping_line, "name", "")
        for item in rows
        if getattr(item, "shipping_line_id", None)
    }
    dates = {item.call_date for item in rows if getattr(item, "call_date", None)}
    pax_total = sum(booking_pax(item, pax_basis=pax_basis) for item in rows)
    kpis = [
        ExportKpi("Arribos", str(len(rows))),
        ExportKpi("PAX", format_pax_count(pax_total)),
        ExportKpi("Puertos", str(len(ports))),
    ]
    if kind == "availability":
        kpis.append(ExportKpi("Días", str(len(dates))))
    else:
        kpis.append(ExportKpi("Navieras", str(len(lines))))
    return kpis


def _criteria_from_query(
    query, *, start, end, pax_basis: str = PAX_BASIS_PLANNED
) -> list[tuple[str, str]]:
    if query is None:
        items: list[tuple[str, str]] = []
        if start and end:
            items.append(("Rango", range_span_label(start, end)))
        if normalize_pax_basis(pax_basis) == PAX_BASIS_CAPACITY:
            items.append(("Base PAX", "Cap. máx."))
        return items
    items: list[tuple[str, str]] = []
    if start and end:
        items.append(("Rango", range_span_label(start, end)))
    ports = _names("catalogs.Port", _ids(query, "port", "ports"))
    items.append(("Puertos", ", ".join(ports) if ports else "todos"))
    vessels = _names("catalogs.Vessel", _ids(query, "vessel"))
    if vessels:
        items.append(("Barco", ", ".join(vessels)))
    lines = _names("catalogs.ShippingLine", _ids(query, "shipping_line"))
    if lines:
        items.append(("Naviera", ", ".join(lines)))
    groups = _names("catalogs.ShippingLineGroup", _ids(query, "shipping_line_group"))
    if groups:
        items.append(("Grupo", ", ".join(groups)))
    position = _position_label(_ids(query, "position"))
    if position:
        items.append(("Posición", position))
    statuses = _status_labels(query)
    if statuses:
        items.append(("Estados", ", ".join(statuses)))
    conflict = _conflict_label(query)
    if conflict:
        items.append(("Conflicto", conflict))
    if _truthy(query, "first_arrival"):
        items.append(("Primer arribo", "sí"))
    reason = _get(query, "cancellation_reason")
    if reason:
        items.append(("Cancelación", _reason_label(reason)))
    search = _get(query, "search")
    if search:
        items.append(("Búsqueda", search))
    sizes = _get(query, "size_category") or _get(query, "size_categories")
    if sizes:
        items.append(("Tamaño", sizes.replace(",", ", ")))
    tags = _names("bookings.BookingTag", _ids(query, "tags"))
    if tags:
        items.append(("Etiqueta", ", ".join(tags)))
    call_dates = _get(query, "call_dates")
    if call_dates:
        count = len([part for part in call_dates.split(",") if part.strip()])
        items.append(("Fechas importadas", str(count)))
    if normalize_pax_basis(pax_basis) == PAX_BASIS_CAPACITY:
        items.append(("Base PAX", "Cap. máx."))
    return items


def _dedupe_criteria(items: list[tuple[str, str]]) -> list[tuple[str, str]]:
    seen: set[str] = set()
    out: list[tuple[str, str]] = []
    for label, value in items:
        if not value or label in seen:
            continue
        seen.add(label)
        out.append((label, value))
    return out


def _ids(query, *keys) -> list[int]:
    raw: list[str] = []
    for key in keys:
        if hasattr(query, "getlist"):
            raw.extend(str(part) for part in (query.getlist(key) or []))
        value = _get(query, key)
        if value:
            raw.append(value)
    out: list[int] = []
    for chunk in raw:
        for part in str(chunk).split(","):
            part = part.strip()
            if not part:
                continue
            try:
                number = int(part)
            except (TypeError, ValueError):
                continue
            if number > 0 and number not in out:
                out.append(number)
    return out


def _get(query, key: str) -> str:
    if query is None:
        return ""
    value = query.get(key)
    return str(value).strip() if value is not None else ""


def _query_date(query, *keys) -> date | None:
    for key in keys:
        raw = _get(query, key)
        if not raw:
            continue
        try:
            return date.fromisoformat(raw)
        except ValueError:
            continue
    return None


def _truthy(query, key: str) -> bool:
    return _get(query, key).lower() in {"1", "true", "yes", "si", "sí"}


def _names(model_label: str, ids: list[int]) -> list[str]:
    if not ids:
        return []
    from django.apps import apps

    try:
        model = apps.get_model(model_label)
    except LookupError:
        return []
    found = {item.id: item for item in model.objects.filter(pk__in=ids)}
    names: list[str] = []
    for pk in ids:
        item = found.get(pk)
        if item is None:
            continue
        names.append(str(getattr(item, "name", None) or getattr(item, "code", pk)))
    return names


def _position_label(ids: list[int]) -> str:
    if not ids:
        return ""
    from apps.catalogs.models import Position
    from apps.catalogs.utils.position_code import position_short_code

    labels: list[str] = []
    for position in Position.objects.filter(pk__in=ids).select_related("port"):
        labels.append(position_short_code(position.port.code, position.code))
    return ", ".join(labels)


def _status_labels(query) -> list[str]:
    from apps.bookings.utils.status_query import parse_status_query_params

    codes = parse_status_query_params(query) if hasattr(query, "getlist") else []
    if not codes:
        raw = _get(query, "status")
        codes = [part.strip().lower() for part in raw.split(",") if part.strip()]
    labels: list[str] = []
    for code in codes:
        labels.append(STATUS_FILTER_LABELS.get(code, code.upper()))
    return labels


def _conflict_label(query) -> str:
    parts: list[str] = []
    flag = _get(query, "has_conflict").lower()
    if flag in {"1", "true", "yes", "si", "sí"}:
        parts.append("Con conflicto")
    elif flag in {"0", "false", "no"}:
        parts.append("Sin conflicto")
    severity = _get(query, "conflict_severity").lower()
    if severity in CONFLICT_LABELS:
        parts.append(CONFLICT_LABELS[severity])
    kind = _get(query, "conflict_type").lower()
    if kind in CONFLICT_LABELS:
        parts.append(CONFLICT_LABELS[kind])
    return " · ".join(parts)


def _reason_label(code: str) -> str:
    from apps.bookings.models import CancellationReason

    try:
        return CancellationReason(code).label
    except ValueError:
        return code
