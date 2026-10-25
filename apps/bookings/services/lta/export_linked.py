"""Export bookings linked to a LongTermAgreement (Excel / CSV)."""

from __future__ import annotations

import csv
import re
from io import BytesIO, StringIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from apps.bookings.services.arrivals_pdf.logo import add_worksheet_port_logo

from apps.bookings.models import Booking, BookingStatus, LongTermAgreement
from apps.bookings.services.booking_export import (
    STATUS_LABELS_ES,
    _format_time,
    vessel_size_label,
)
from apps.catalogs.utils.position_code import position_short_code

HEADERS = [
    "Código",
    "Estado",
    "Puerto",
    "Naviera",
    "Barco",
    "Tamaño",
    "Posición",
    "Fecha de escala",
    "ETA",
    "ETD",
    "Tipo de regla",
    "Acuerdo",
]


def _slug_filename(value: str) -> str:
    text = re.sub(r"[^\w\-]+", "-", (value or "").strip(), flags=re.UNICODE)
    text = re.sub(r"-{2,}", "-", text).strip("-").lower()
    return text[:80] or "acuerdo"


def _extra_dates(agreement: LongTermAgreement) -> set[str]:
    out: set[str] = set()
    for item in agreement.date_exceptions or []:
        if not isinstance(item, dict):
            continue
        if str(item.get("kind") or "").lower() != "include":
            continue
        raw = item.get("date")
        if raw:
            out.add(str(raw)[:10])
    return out


def _rule_type(booking: Booking, extras: set[str]) -> str:
    call = booking.call_date.isoformat() if booking.call_date else ""
    return "Extra" if call in extras else "Según acuerdo"


def _position_label(booking: Booking) -> str:
    if not booking.position_id or not booking.position:
        return ""
    port_code = ""
    if booking.port_id and booking.port:
        port_code = booking.port.code or ""
    if port_code:
        return position_short_code(port_code, booking.position.code)
    return booking.position.code


def _row(booking: Booking, agreement: LongTermAgreement, extras: set[str]) -> list:
    return [
        booking.booking_code,
        STATUS_LABELS_ES.get(booking.status, booking.get_status_display()),
        booking.port.name if booking.port_id else "",
        booking.shipping_line.name if booking.shipping_line_id else "",
        booking.vessel.name if booking.vessel_id else "",
        vessel_size_label(booking),
        _position_label(booking),
        booking.call_date.isoformat() if booking.call_date else "",
        _format_time(booking.eta),
        _format_time(booking.etd),
        _rule_type(booking, extras),
        agreement.code,
    ]


def linked_bookings_queryset(agreement: LongTermAgreement):
    return (
        Booking.objects.filter(long_term_agreement_id=agreement.pk)
        .exclude(status=BookingStatus.C)
        .select_related("port", "shipping_line", "vessel", "position")
        .order_by("call_date", "id")
    )


def export_linked_filename(agreement: LongTermAgreement, fmt: str) -> str:
    stamp = timezone.localdate().isoformat()
    return f"lta-{_slug_filename(agreement.code)}-reservas-{stamp}.{fmt}"


def build_linked_bookings_xlsx(agreement: LongTermAgreement) -> bytes:
    extras = _extra_dates(agreement)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reservas del acuerdo"
    placed = 0
    if add_worksheet_port_logo(
        sheet, agreement.port, row=1, col=1, max_w=88, max_h=36
    ):
        placed = 1
    title_start = 3 if placed else 1
    sheet.merge_cells(
        start_row=1,
        start_column=title_start,
        end_row=1,
        end_column=len(HEADERS),
    )
    title = sheet.cell(1, title_start, "Reservas del acuerdo")
    title.font = Font(name="Calibri", size=16, bold=True)
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 40 if placed else 22
    sheet.append(HEADERS)
    for cell in sheet[2]:
        cell.font = Font(bold=True)
    for booking in linked_bookings_queryset(agreement).iterator(chunk_size=200):
        sheet.append(_row(booking, agreement, extras))
    for index, column_cells in enumerate(sheet.columns, start=1):
        max_len = 0
        for cell in column_cells:
            if isinstance(cell, MergedCell) or cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        sheet.column_dimensions[get_column_letter(index)].width = min(max_len + 2, 48)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_linked_bookings_csv(agreement: LongTermAgreement) -> bytes:
    extras = _extra_dates(agreement)
    buffer = StringIO()
    writer = csv.writer(buffer)
    writer.writerow(HEADERS)
    for booking in linked_bookings_queryset(agreement).iterator(chunk_size=200):
        writer.writerow(_row(booking, agreement, extras))
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")
