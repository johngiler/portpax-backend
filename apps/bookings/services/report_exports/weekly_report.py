"""Reporte Semanal — DESGLOSE DE MOVIMIENTOS (ops week × call year × port)."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.audit.models import BookingAuditEntry
from apps.bookings.models import Booking, BookingStatus
from apps.bookings.services.report_exports.booking_movements import (
    MIN_REPORT_YEAR,
    _classify_kinds,
    _resolve_call_year,
    _resolve_port_id,
    parse_movement_year,
)
from apps.bookings.services.report_exports.common import (
    booking_pax,
    normalize_pax_basis,
    pax_basis_note,
    scheduled_bookings_qs,
    PAX_BASIS_CAPACITY,
    PAX_BASIS_PLANNED,
)
from apps.bookings.services.report_exports.report_theme import (
    NAVY,
    TEXT,
    WHITE,
)
from apps.bookings.services.report_exports.xlsx_style import (
    ALIGN_CENTER,
    ALIGN_LEFT,
    BORDER_NONE,
    prepare_report_sheet,
    style_cell,
)
from apps.bookings.services.validation.legend_labels import port_legend_label
from apps.catalogs.models import Port

# Internal kind → export/UI label (image layout).
WEEKLY_METRICS: tuple[tuple[str, str], ...] = (
    ("NEW BOOKING", "NEW BOOKING"),
    ("CANCELLATION", "CANCELLATION"),
    ("SHIP CHANGE", "SHIP CHANGE"),
    ("REAL PAX", "PAX PROY / REAL"),
)

WEEKLY_KIND_KEYS = tuple(k for k, _ in WEEKLY_METRICS)

_WEEK_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
_PORT_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
_PORT_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
_METRIC_FONT = Font(name="Calibri", size=10, italic=True, color="5B9BD5")
_YEAR_FONT = Font(name="Calibri", size=11, bold=True, color="2F5496")
_TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="2F5496")
_THIN = Side(style="thin", color="2F5496")


def max_iso_week(year: int) -> int:
    """ISO weeks in ``year`` (52 or 53)."""
    return date(year, 12, 28).isocalendar()[1]


def iso_week_bounds(year: int, week: int) -> tuple[date, date]:
    """Monday–Sunday dates for ISO week ``week`` of ISO ``year``."""
    jan4 = date(year, 1, 4)
    week1_monday = jan4 - timedelta(days=jan4.isoweekday() - 1)
    start = week1_monday + timedelta(weeks=week - 1)
    return start, start + timedelta(days=6)


def current_iso_year_week() -> tuple[int, int]:
    today = timezone.localdate()
    iso = today.isocalendar()
    return int(iso.year), int(iso.week)


def parse_movement_week(raw: str | None, *, year: int) -> int:
    """Week 1…max for ``year``; default = current ISO week when year matches."""
    max_w = max_iso_week(year)
    cur_y, cur_w = current_iso_year_week()
    default = cur_w if year == cur_y else 1
    default = max(1, min(default, max_w))
    if raw is None or str(raw).strip() == "":
        return default
    try:
        week = int(str(raw).strip().split(",")[0])
    except ValueError:
        return default
    if week < 1:
        return 1
    if week > max_w:
        return max_w
    return week


def parse_weekly_year(raw: str | None) -> int:
    """Same ops-year window as Movimientos de bookings (past + current)."""
    from apps.bookings.services.report_exports.booking_movements import (
        parse_movement_year,
    )

    return parse_movement_year(raw)


def port_report_abbrev(port: Port | None) -> str:
    """Legacy short code helper (exports no longer show abbrevs in labels)."""
    if port is None:
        return ""
    code = (port.code or "").strip().lower().replace("_", "").replace("-", "")
    if len(code) >= 3:
        return code[:3].upper()
    name = (getattr(port, "name", None) or "").strip()
    return (name[:3] or "?").upper()


def _media_url(request, field) -> str | None:
    if not field:
        return None
    try:
        url = field.url
    except ValueError:
        return None
    if request is not None:
        return request.build_absolute_uri(url)
    return url


def _projected_capacity(
    booking: Booking | None,
    *,
    pax_basis: str = PAX_BASIS_PLANNED,
) -> int:
    """Projected figure for PAX PROY / REAL (respects Base PAX filter)."""
    if booking is None:
        return 0
    basis = normalize_pax_basis(pax_basis)
    vessel = getattr(booking, "vessel", None)
    cap = getattr(vessel, "pax_capacity", None) if vessel is not None else None
    planned = int(booking.planned_pax) if booking.planned_pax is not None else None
    if basis == PAX_BASIS_CAPACITY:
        if cap is not None:
            return int(cap)
        return planned or 0
    if planned is not None:
        return planned
    return int(cap) if cap is not None else 0


def _pax_for_kind(
    entry: BookingAuditEntry,
    kind: str,
    booking: Booking | None,
    *,
    pax_basis: str = PAX_BASIS_PLANNED,
) -> int:
    """Signed PAX attributed to one weekly metric kind."""
    changes = entry.changes or {}
    if kind == "CANCELLATION":
        base = booking_pax(booking, pax_basis=pax_basis) if booking else 0
        return -int(base)
    if kind == "REAL PAX":
        # Beto: projected (per Base PAX) − actual manifested PAX of the arrival.
        if booking is None:
            return 0
        projected = _projected_capacity(booking, pax_basis=pax_basis)
        actual: int | None = None
        if "actual_pax" in changes:
            to = (changes.get("actual_pax") or {}).get("to")
            if to is not None:
                try:
                    actual = int(to)
                except (TypeError, ValueError):
                    actual = None
        if actual is None and booking.actual_pax is not None:
            actual = int(booking.actual_pax)
        if actual is None:
            return 0
        return projected - actual
    if kind in ("NEW BOOKING", "SHIP CHANGE"):
        if booking is None:
            return 0
        return int(booking_pax(booking, pax_basis=pax_basis))
    return 0


def _port_year_totals(
    *,
    call_years: list[int],
    allowed: set[int] | None,
    without_lta: bool,
    pax_basis: str,
) -> dict[int, list[int]]:
    """Full-year port PAX totals (not week movements) for each call year column."""
    if not call_years:
        return {}
    date_from = date(call_years[0], 1, 1)
    date_to = date(call_years[-1], 12, 31)
    year_index = {y: i for i, y in enumerate(call_years)}
    totals: dict[int, list[int]] = defaultdict(lambda: [0] * len(call_years))
    qs = scheduled_bookings_qs(
        date_from=date_from,
        date_to=date_to,
        allowed_ports=allowed,
        without_lta=without_lta,
    ).select_related("vessel")
    for booking in qs.iterator(chunk_size=500):
        cy = booking.call_date.year if booking.call_date else None
        if cy is None or cy not in year_index:
            continue
        totals[booking.port_id][year_index[cy]] += booking_pax(
            booking, pax_basis=pax_basis
        )
    return totals


def build_weekly_report(
    *,
    year: int,
    week: int,
    without_lta: bool = False,
    pax_basis: str = PAX_BASIS_PLANNED,
    allowed_ports: set[int] | list[int] | None = None,
    request=None,
) -> dict[str, Any]:
    """
    Ops activity in ISO week ``week`` of ``year``.

    Columns = call years ``year``…``year+3``.
    Blue port row = full-year port PAX totals (all years in the window).
    Metric rows = signed PAX movements registered that week.
    PAX PROY / REAL = projected (Base PAX) − actual_pax for Real updates.
    """
    if year < MIN_REPORT_YEAR:
        raise ValueError(f"year debe ser >= {MIN_REPORT_YEAR}.")
    max_w = max_iso_week(year)
    if week < 1 or week > max_w:
        raise ValueError(f"week debe estar entre 1 y {max_w} para {year}.")

    basis = normalize_pax_basis(pax_basis)

    allowed: set[int] | None
    if allowed_ports is None:
        allowed = None
    else:
        allowed = {int(x) for x in allowed_ports}

    week_start, week_end = iso_week_bounds(year, week)
    call_years = [year, year + 1, year + 2, year + 3]
    year_index = {y: i for i, y in enumerate(call_years)}

    tz = timezone.get_current_timezone()
    start_dt = timezone.make_aware(
        datetime(week_start.year, week_start.month, week_start.day, 0, 0, 0),
        tz,
    )
    end_dt = timezone.make_aware(
        datetime(week_end.year, week_end.month, week_end.day, 23, 59, 59),
        tz,
    )

    entries = list(
        BookingAuditEntry.objects.filter(
            created_at__gte=start_dt,
            created_at__lte=end_dt,
        )
        .order_by("created_at", "id")
    )
    booking_ids = {e.booking_id for e in entries if e.booking_id}
    bookings = {
        b.pk: b
        for b in Booking.objects.filter(pk__in=booking_ids).select_related(
            "port",
            "vessel",
        )
    }

    # port_id -> kind -> [pax per call_year]
    pax_map: dict[int, dict[str, list[int]]] = defaultdict(
        lambda: {k: [0] * len(call_years) for k in WEEKLY_KIND_KEYS}
    )

    for entry in entries:
        kinds = _classify_kinds(entry)
        if not kinds:
            continue
        booking = bookings.get(entry.booking_id) if entry.booking_id else None
        if (
            without_lta
            and booking is not None
            and booking.status == BookingStatus.LTA
        ):
            continue
        port_id = _resolve_port_id(entry, booking)
        if port_id is None:
            continue
        if allowed is not None and port_id not in allowed:
            continue
        call_year = _resolve_call_year(entry, booking)
        if call_year is None or call_year not in year_index:
            continue
        yi = year_index[call_year]
        for kind in kinds:
            if kind not in WEEKLY_KIND_KEYS:
                continue
            delta = _pax_for_kind(entry, kind, booking, pax_basis=basis)
            if delta == 0:
                continue
            pax_map[port_id][kind][yi] += delta

    port_year_totals = _port_year_totals(
        call_years=call_years,
        allowed=allowed,
        without_lta=without_lta,
        pax_basis=basis,
    )

    port_ids = sorted(set(pax_map.keys()) | set(port_year_totals.keys()))
    ports = {
        p.pk: p
        for p in Port.objects.filter(pk__in=port_ids).only(
            "id", "code", "name", "commercial_name", "logo"
        )
    }
    port_ids = sorted(
        port_ids,
        key=lambda pid: (
            port_legend_label(ports.get(pid)).lower(),
            pid,
        ),
    )

    port_rows: list[dict[str, Any]] = []
    for port_id in port_ids:
        port = ports.get(port_id)
        kind_map = pax_map.get(port_id) or {
            k: [0] * len(call_years) for k in WEEKLY_KIND_KEYS
        }
        metrics: list[dict[str, Any]] = []
        any_movement = False
        for key, label in WEEKLY_METRICS:
            values = list(kind_map.get(key) or [0] * len(call_years))
            if any(values):
                any_movement = True
            metrics.append({"key": key, "label": label, "values": values})
        totals = list(
            port_year_totals.get(port_id) or [0] * len(call_years)
        )
        if not any_movement and not any(totals):
            continue
        port_rows.append(
            {
                "port_id": port_id,
                "port_name": port_legend_label(port) if port else f"Puerto #{port_id}",
                "logo": _media_url(request, port.logo) if port else None,
                "totals": totals,
                "metrics": metrics,
            }
        )

    note = (
        "Fila azul = totales PAX del puerto por año de escala (no la suma de la semana). "
        "Filas inferiores = movimientos registrados en la semana ISO. "
        "PAX PROY / REAL = proyectado (Base PAX) − PAX real del arribo. "
        f"{pax_basis_note(basis)}"
    )
    if without_lta:
        note = f"{note} Sin LTA."

    return {
        "kind": "weekly_report",
        "title": "Reporte Semanal",
        "report_name": "Reporte Semanal",
        "year": year,
        "week": week,
        "week_start": week_start.isoformat(),
        "week_end": week_end.isoformat(),
        "without_lta": without_lta,
        "pax_basis": basis,
        "call_years": call_years,
        "metric_labels": [label for _, label in WEEKLY_METRICS],
        "ports": port_rows,
        "note": note,
    }


def weekly_report_filename(year: int, week: int, ext: str = "xlsx") -> str:
    _ = (year, week)
    from apps.bookings.services.report_exports.filenames import report_download_filename

    return report_download_filename("weekly_report", ext)


def _fmt_num(n: int) -> str | int:
    if not n:
        return ""
    return n


def build_weekly_report_xlsx(payload: dict[str, Any]) -> bytes:
    call_years: list[int] = list(payload.get("call_years") or [])
    ports: list[dict[str, Any]] = list(payload.get("ports") or [])
    col_count = 1 + len(call_years)

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte Semanal"
    prepare_report_sheet(ws)

    # Title + week badge row.
    title = str(payload.get("title") or payload.get("report_name") or "Reporte Semanal")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, col_count - 1))
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = _TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.border = Border(
        left=_THIN, right=_THIN, top=_THIN, bottom=_THIN
    )
    for c in range(1, max(2, col_count - 1) + 1):
        ws.cell(row=1, column=c).border = Border(
            left=_THIN if c == 1 else None,
            right=_THIN if c == max(2, col_count - 1) else None,
            top=_THIN,
            bottom=_THIN,
        )

    week_col = col_count
    week_cell = ws.cell(row=1, column=week_col, value=f"Sem.\n{payload.get('week')}")
    week_cell.fill = _WEEK_FILL
    week_cell.font = Font(name="Calibri", size=11, bold=True, color=TEXT)
    week_cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    ws.row_dimensions[1].height = 36

    # Column headers
    header_row = 3
    style_cell(
        ws.cell(row=header_row, column=1, value="PUERTO"),
        font=Font(name="Calibri", size=10, italic=True, color="2F5496"),
        fill=None,
        alignment=ALIGN_LEFT,
        border=BORDER_NONE,
    )
    for i, y in enumerate(call_years):
        style_cell(
            ws.cell(row=header_row, column=2 + i, value=y),
            font=_YEAR_FONT,
            fill=None,
            alignment=ALIGN_CENTER,
            border=BORDER_NONE,
        )

    row = header_row + 1
    for port in ports:
        label = str(port.get("port_name") or "").strip()
        cell = ws.cell(row=row, column=1, value=label)
        cell.fill = _PORT_FILL
        cell.font = _PORT_FONT
        cell.alignment = ALIGN_LEFT
        cell.border = BORDER_NONE
        totals = list(port.get("totals") or [])
        for i, y in enumerate(call_years):
            val = int(totals[i] if i < len(totals) else 0)
            c = ws.cell(row=row, column=2 + i, value=val)
            c.fill = _PORT_FILL
            c.font = _PORT_FONT
            c.alignment = ALIGN_CENTER
            c.border = BORDER_NONE
            c.number_format = "#,##0"
        row += 1
        for metric in port.get("metrics") or []:
            mcell = ws.cell(
                row=row, column=1, value=str(metric.get("label") or "")
            )
            mcell.font = _METRIC_FONT
            mcell.alignment = ALIGN_LEFT
            mcell.border = BORDER_NONE
            values = list(metric.get("values") or [])
            for i, _y in enumerate(call_years):
                val = values[i] if i < len(values) else 0
                c = ws.cell(row=row, column=2 + i, value=_fmt_num(int(val or 0)))
                c.font = Font(name="Calibri", size=10, color=TEXT)
                c.alignment = ALIGN_CENTER
                c.border = BORDER_NONE
                if isinstance(c.value, int):
                    c.number_format = "#,##0;-#,##0;"
            row += 1
        row += 1  # gap between ports

    ws.column_dimensions["A"].width = 28
    for i in range(len(call_years)):
        ws.column_dimensions[get_column_letter(2 + i)].width = 14

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_weekly_report_csv(payload: dict[str, Any]) -> bytes:
    import csv
    from io import StringIO

    call_years: list[int] = list(payload.get("call_years") or [])
    buf = StringIO()
    writer = csv.writer(buf)
    writer.writerow([payload.get("title") or "Reporte Semanal"])
    writer.writerow(
        [
            f"Sem. {payload.get('week')}",
            f"{payload.get('week_start')} → {payload.get('week_end')}",
        ]
    )
    writer.writerow([])
    writer.writerow(["PUERTO", *[str(y) for y in call_years]])
    for port in payload.get("ports") or []:
        label = str(port.get("port_name") or "").strip()
        writer.writerow([label, *[port.get("totals") or [0] * len(call_years)]])
        for metric in port.get("metrics") or []:
            writer.writerow(
                [
                    metric.get("label") or "",
                    *(metric.get("values") or [0] * len(call_years)),
                ]
            )
        writer.writerow([])
    return buf.getvalue().encode("utf-8-sig")
