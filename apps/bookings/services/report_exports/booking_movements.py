"""Movimientos de bookings — audit event counts + signed PAX by port/call year."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from typing import Any

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Border, Font, Side
from openpyxl.utils import get_column_letter

from apps.audit.models import BookingAuditEntry
from apps.bookings.models import Booking, BookingStatus
from apps.bookings.services.report_exports.common import booking_pax
from apps.bookings.services.report_exports.xlsx_style import (
    ALIGN_CENTER,
    ALIGN_LEFT,
    ALIGN_RIGHT,
    FILL_ALT,
    FILL_HEADER,
    FILL_TITLE,
    FILL_TOTAL,
    FONT_DATA,
    FONT_HEADER,
    FONT_NOTE,
    FONT_TOTAL,
    NAVY,
    style_cell,
    write_title_row,
)
from apps.bookings.services.validation.legend_labels import port_legend_label
from apps.catalogs.models import Port

MIN_REPORT_YEAR = 2025

MOVEMENT_KINDS = (
    "CANCELLATION",
    "DATE CHANGE",
    "NEW BOOKING",
    "REAL PAX",
    "SHIP CHANGE",
)

MONTH_LABELS_ES = (
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
)

_SKIP_ACTIONS = frozenset(
    {
        "conflict_detected",
        "conflict_updated",
        "conflict_resolved",
        "lta_linked",
        "lta_unlinked",
        "deleted",
    }
)

_BLACK = Side(style="thin", color="000000")
BORDER_BLACK = Border(left=_BLACK, right=_BLACK, top=_BLACK, bottom=_BLACK)


def parse_movement_year(raw: str | None) -> int:
    """Single ops year for the report (min 2025)."""
    now_y = timezone.localdate().year
    default = max(MIN_REPORT_YEAR, min(now_y, now_y + 4))
    if raw is None or str(raw).strip() == "":
        return default
    try:
        year = int(str(raw).strip().split(",")[0])
    except ValueError:
        return default
    if year < MIN_REPORT_YEAR:
        return MIN_REPORT_YEAR
    return year


def _parse_iso_date(raw: Any) -> date | None:
    if raw is None or raw == "":
        return None
    if isinstance(raw, date) and not isinstance(raw, datetime):
        return raw
    if isinstance(raw, datetime):
        return raw.date()
    text = str(raw).strip()[:10]
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def _classify_kinds(entry: BookingAuditEntry) -> list[str]:
    action = entry.action or ""
    changes = entry.changes or {}
    if action in _SKIP_ACTIONS:
        return []
    if action == "created":
        return ["NEW BOOKING"]
    if action == "status_change":
        to = (changes.get("status") or {}).get("to")
        if to == BookingStatus.C:
            return ["CANCELLATION"]
        if to == BookingStatus.R:
            return ["REAL PAX"]
        return []
    if action == "identity_update":
        kinds: list[str] = []
        if "call_date" in changes:
            kinds.append("DATE CHANGE")
        if "vessel_id" in changes or "vessel" in changes:
            kinds.append("SHIP CHANGE")
        return kinds
    if action == "operational_update" and "actual_pax" in changes:
        return ["REAL PAX"]
    return []


def _pax_delta_for_block_b(
    entry: BookingAuditEntry,
    kinds: list[str],
    booking: Booking | None,
) -> int:
    """Signed PAX contribution once per audit entry (not per kind)."""
    if not kinds:
        return 0
    changes = entry.changes or {}
    if "CANCELLATION" in kinds:
        base = booking_pax(booking, pax_basis="planned") if booking else 0
        return -int(base)
    if kinds == ["REAL PAX"] and "actual_pax" in changes:
        ch = changes.get("actual_pax") or {}
        fr, to = ch.get("from"), ch.get("to")
        if fr is not None or to is not None:
            try:
                return int(to or 0) - int(fr or 0)
            except (TypeError, ValueError):
                pass
    if booking is None:
        return 0
    return int(booking_pax(booking, pax_basis="planned"))


def _resolve_port_id(
    entry: BookingAuditEntry,
    booking: Booking | None,
) -> int | None:
    if entry.port_id:
        return int(entry.port_id)
    if booking is not None and booking.port_id:
        return int(booking.port_id)
    entity = (entry.changes or {}).get("entity") or {}
    raw = entity.get("port_id")
    if raw is None:
        return None
    try:
        return int(raw)
    except (TypeError, ValueError):
        return None


def _resolve_call_year(
    entry: BookingAuditEntry,
    booking: Booking | None,
) -> int | None:
    if booking is not None and booking.call_date:
        return booking.call_date.year
    entity = (entry.changes or {}).get("entity") or {}
    call = _parse_iso_date(entity.get("call_date"))
    if call:
        return call.year
    changes = entry.changes or {}
    if "call_date" in changes:
        call = _parse_iso_date((changes.get("call_date") or {}).get("to"))
        if call:
            return call.year
    return None


def build_booking_movements_report(
    *,
    year: int,
    allowed_ports: set[int] | list[int] | None = None,
) -> dict[str, Any]:
    """
    Block A: event counts by movement kind × ops month.
    Block B: signed PAX by port × call year × ops month.
    """
    if year < MIN_REPORT_YEAR:
        raise ValueError(f"year debe ser >= {MIN_REPORT_YEAR}.")

    allowed: set[int] | None
    if allowed_ports is None:
        allowed = None
    else:
        allowed = {int(x) for x in allowed_ports}

    tz = timezone.get_current_timezone()
    start = timezone.make_aware(datetime(year, 1, 1, 0, 0, 0), tz)
    end = timezone.make_aware(datetime(year, 12, 31, 23, 59, 59), tz)

    entries = list(
        BookingAuditEntry.objects.filter(
            created_at__gte=start,
            created_at__lte=end,
        )
        .exclude(action__in=_SKIP_ACTIONS)
        .order_by("created_at", "id")
    )
    booking_ids = {e.booking_id for e in entries if e.booking_id}
    bookings = {
        b.pk: b
        for b in Booking.objects.filter(pk__in=booking_ids).select_related(
            "port",
            "vessel",
        )
    }

    counts: dict[str, list[int]] = {k: [0] * 12 for k in MOVEMENT_KINDS}
    # port_id -> call_year -> month(0-11) -> pax
    pax_map: dict[int, dict[int, list[int]]] = defaultdict(
        lambda: defaultdict(lambda: [0] * 12)
    )

    for entry in entries:
        kinds = _classify_kinds(entry)
        if not kinds:
            continue
        local_dt = timezone.localtime(entry.created_at, tz)
        month_idx = local_dt.month - 1
        for kind in kinds:
            counts[kind][month_idx] += 1

        booking = bookings.get(entry.booking_id) if entry.booking_id else None
        port_id = _resolve_port_id(entry, booking)
        if port_id is None:
            continue
        if allowed is not None and port_id not in allowed:
            continue
        call_year = _resolve_call_year(entry, booking)
        if call_year is None:
            continue
        delta = _pax_delta_for_block_b(entry, kinds, booking)
        if delta == 0:
            continue
        pax_map[port_id][call_year][month_idx] += delta

    type_rows: list[dict[str, Any]] = []
    month_totals_a = [0] * 12
    for kind in MOVEMENT_KINDS:
        months = counts[kind]
        total = sum(months)
        type_rows.append(
            {
                "kind": kind,
                "months": months,
                "total": total,
            }
        )
        for i, v in enumerate(months):
            month_totals_a[i] += v
    grand_a = sum(month_totals_a)
    pct_a = [
        int(round((v / grand_a) * 100)) if grand_a else 0 for v in month_totals_a
    ]
    avg_a = int(round(grand_a / 12)) if grand_a else 0

    port_ids = sorted(pax_map.keys())
    ports = {
        p.pk: p
        for p in Port.objects.filter(pk__in=port_ids).order_by("name", "code")
    }
    # Stable order by friendly name
    port_ids = sorted(
        port_ids,
        key=lambda pid: (
            port_legend_label(ports.get(pid)).lower(),
            pid,
        ),
    )

    port_blocks: list[dict[str, Any]] = []
    month_totals_b = [0] * 12
    for port_id in port_ids:
        port = ports.get(port_id)
        label = port_legend_label(port) if port else f"Puerto #{port_id}"
        year_map = pax_map[port_id]
        call_years = sorted(y for y, months in year_map.items() if any(months))
        if not call_years:
            continue
        year_rows: list[dict[str, Any]] = []
        port_months = [0] * 12
        for cy in call_years:
            months = list(year_map[cy])
            year_rows.append(
                {
                    "year": cy,
                    "months": months,
                    "total": sum(months),
                }
            )
            for i, v in enumerate(months):
                port_months[i] += v
        for i, v in enumerate(port_months):
            month_totals_b[i] += v
        port_blocks.append(
            {
                "port_id": port_id,
                "port_name": label,
                "months": port_months,
                "total": sum(port_months),
                "years": year_rows,
            }
        )

    grand_b = sum(month_totals_b)
    pct_b = [
        int(round((v / grand_b) * 100)) if grand_b else 0 for v in month_totals_b
    ]
    avg_b = int(round(grand_b / 12)) if grand_b else 0

    return {
        "kind": "booking_movements",
        "title": "Movimientos de bookings",
        "year": year,
        "month_labels": list(MONTH_LABELS_ES),
        "type_rows": type_rows,
        "type_month_totals": month_totals_a,
        "type_grand_total": grand_a,
        "type_month_pct": pct_a,
        "type_avg_per_month": avg_a,
        "port_blocks": port_blocks,
        "pax_month_totals": month_totals_b,
        "pax_grand_total": grand_b,
        "pax_month_pct": pct_b,
        "pax_avg_per_month": avg_b,
        "note": (
            "Conteos de movimientos por mes de registro. "
            "PAX con signo (+ altas/cambios, − cancelaciones) por puerto y año de escala."
        ),
    }


def booking_movements_filename(year: int) -> str:
    _ = year
    return "Movimientos de bookings.xlsx"


def _write_matrix_header(ws, row: int, col_span_start: int = 1) -> None:
    cell = ws.cell(row=row, column=1, value="")
    style_cell(
        cell,
        font=FONT_HEADER,
        fill=FILL_HEADER,
        alignment=ALIGN_LEFT,
        border=BORDER_BLACK,
    )
    for i, label in enumerate(MONTH_LABELS_ES, start=2):
        cell = ws.cell(row=row, column=i, value=label)
        style_cell(
            cell,
            font=FONT_HEADER,
            fill=FILL_HEADER,
            alignment=ALIGN_CENTER,
            border=BORDER_BLACK,
        )
    cell = ws.cell(row=row, column=14, value="Total general")
    style_cell(
        cell,
        font=FONT_HEADER,
        fill=FILL_HEADER,
        alignment=ALIGN_CENTER,
        border=BORDER_BLACK,
    )


def _write_data_row(
    ws,
    row: int,
    *,
    label: str,
    months: list[int],
    total: int,
    bold: bool = False,
    fill=None,
) -> None:
    font = FONT_TOTAL if bold else FONT_DATA
    row_fill = fill or FILL_ALT
    cell = ws.cell(row=row, column=1, value=label)
    style_cell(
        cell,
        font=font,
        fill=row_fill,
        alignment=ALIGN_LEFT,
        border=BORDER_BLACK,
    )
    for i, value in enumerate(months, start=2):
        cell = ws.cell(row=row, column=i, value=value if value else "")
        style_cell(
            cell,
            font=font,
            fill=row_fill,
            alignment=ALIGN_RIGHT,
            border=BORDER_BLACK,
            number_format="#,##0",
        )
    cell = ws.cell(row=row, column=14, value=total if total else "")
    style_cell(
        cell,
        font=font,
        fill=FILL_TOTAL if bold else row_fill,
        alignment=ALIGN_RIGHT,
        border=BORDER_BLACK,
        number_format="#,##0",
    )


def _write_pct_row(ws, row: int, pcts: list[int]) -> None:
    cell = ws.cell(row=row, column=1, value="")
    style_cell(cell, font=FONT_NOTE, fill=FILL_ALT, border=BORDER_BLACK)
    for i, pct in enumerate(pcts, start=2):
        cell = ws.cell(row=row, column=i, value=f"{pct}%" if pct else "")
        style_cell(
            cell,
            font=FONT_NOTE,
            fill=FILL_ALT,
            alignment=ALIGN_CENTER,
            border=BORDER_BLACK,
        )
    cell = ws.cell(row=row, column=14, value="")
    style_cell(cell, font=FONT_NOTE, fill=FILL_ALT, border=BORDER_BLACK)


def build_booking_movements_xlsx(payload: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"

    year = int(payload.get("year") or 0)
    write_title_row(ws, 1, str(payload.get("title") or "Movimientos de bookings"), 14)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=14)
    cell = ws.cell(row=2, column=1, value=f"Año {year}" if year else "")
    style_cell(
        cell,
        font=Font(name="Calibri", size=12, bold=True, color=NAVY),
        fill=FILL_TITLE,
        alignment=ALIGN_LEFT,
        border=None,
    )
    note = payload.get("note") or ""
    if note:
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=14)
        cell = ws.cell(row=3, column=1, value=note)
        style_cell(cell, font=FONT_NOTE, alignment=ALIGN_LEFT, border=None)
        row = 5
    else:
        row = 4

    # Block A
    _write_matrix_header(ws, row)
    row += 1
    for item in payload.get("type_rows") or []:
        _write_data_row(
            ws,
            row,
            label=str(item.get("kind") or ""),
            months=list(item.get("months") or [0] * 12),
            total=int(item.get("total") or 0),
        )
        row += 1
    _write_data_row(
        ws,
        row,
        label="Total general",
        months=list(payload.get("type_month_totals") or [0] * 12),
        total=int(payload.get("type_grand_total") or 0),
        bold=True,
        fill=FILL_TOTAL,
    )
    row += 1
    _write_pct_row(ws, row, list(payload.get("type_month_pct") or [0] * 12))
    row += 1
    cell = ws.cell(
        row=row,
        column=16,
        value=f"Promedio por mes: {int(payload.get('type_avg_per_month') or 0):,}",
    )
    style_cell(cell, font=FONT_TOTAL, alignment=ALIGN_LEFT, border=None)
    row += 2

    # Block B
    _write_matrix_header(ws, row)
    row += 1
    for block in payload.get("port_blocks") or []:
        _write_data_row(
            ws,
            row,
            label=str(block.get("port_name") or ""),
            months=list(block.get("months") or [0] * 12),
            total=int(block.get("total") or 0),
            bold=True,
            fill=FILL_TOTAL,
        )
        row += 1
        for year_row in block.get("years") or []:
            _write_data_row(
                ws,
                row,
                label=str(year_row.get("year") or ""),
                months=list(year_row.get("months") or [0] * 12),
                total=int(year_row.get("total") or 0),
            )
            row += 1
    _write_data_row(
        ws,
        row,
        label="Total general",
        months=list(payload.get("pax_month_totals") or [0] * 12),
        total=int(payload.get("pax_grand_total") or 0),
        bold=True,
        fill=FILL_TOTAL,
    )
    row += 1
    _write_pct_row(ws, row, list(payload.get("pax_month_pct") or [0] * 12))
    row += 1
    cell = ws.cell(
        row=row,
        column=16,
        value=f"Promedio por mes: {int(payload.get('pax_avg_per_month') or 0):,}",
    )
    style_cell(cell, font=FONT_TOTAL, alignment=ALIGN_LEFT, border=None)

    ws.column_dimensions["A"].width = 22
    for col in range(2, 15):
        ws.column_dimensions[get_column_letter(col)].width = 11
    ws.column_dimensions["P"].width = 28

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
