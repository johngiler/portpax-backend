"""Resumen de movimientos — year blocks + port/carrier PAX summaries."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, time
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Border, Font, Side
from openpyxl.utils import get_column_letter

from apps.bookings.models import BookingTag
from apps.bookings.services.report_exports.common import (
    booking_pax,
    normalize_pax_basis,
    pax_basis_note,
    scheduled_bookings_qs,
    years_in_range,
)
from apps.bookings.services.report_exports.xlsx_style import (
    ALIGN_CENTER,
    ALIGN_LEFT,
    ALIGN_RIGHT,
    FILL_ALT,
    FILL_HEADER,
    FILL_TITLE,
    FILL_TOTAL,
    FONT_DATA,
    FONT_HEADER,
    FONT_NOTE,
    FONT_TITLE,
    FONT_TOTAL,
    NAVY,
    style_cell,
    write_title_row,
)
from apps.catalogs.models import Port, ShippingLine

MIN_REPORT_YEAR = 2025

_BLACK = Side(style="thin", color="000000")
BORDER_BLACK = Border(left=_BLACK, right=_BLACK, top=_BLACK, bottom=_BLACK)


def _media_url(request, field) -> str | None:
    if not field:
        return None
    try:
        url = field.url
    except ValueError:
        return None
    if request is not None:
        return request.build_absolute_uri(url)
    return url



def parse_id_list(raw: str | None) -> list[int]:
    if not raw:
        return []
    out: list[int] = []
    for part in str(raw).split(","):
        part = part.strip()
        if not part:
            continue
        try:
            value = int(part)
        except ValueError:
            continue
        if value > 0 and value not in out:
            out.append(value)
    return out


def parse_report_years(raw: str | None, *, date_from: date, date_to: date) -> list[int]:
    """Selected calendar years for left blocks (min 2025)."""
    allowed = [
        y
        for y in years_in_range(date_from, date_to)
        if y >= MIN_REPORT_YEAR
    ]
    parsed = parse_id_list(raw)
    if not parsed:
        return allowed
    allowed_set = set(allowed)
    return [y for y in parsed if y in allowed_set]


def _format_time(value: time | None) -> str:
    if value is None:
        return "—"
    return value.strftime("%H:%M")


def _format_arrival(d: date) -> str:
    # 14-Oct-28
    months = (
        "Jan",
        "Feb",
        "Mar",
        "Apr",
        "May",
        "Jun",
        "Jul",
        "Aug",
        "Sep",
        "Oct",
        "Nov",
        "Dec",
    )
    return f"{d.day:02d}-{months[d.month - 1]}-{d.year % 100:02d}"


def _port_label(port: Port) -> str:
    return (port.name or port.code or f"#{port.pk}").strip()


def _pct(carrier_pax: int, port_pax: int) -> str:
    if port_pax <= 0:
        return "—"
    return f"{int(round((carrier_pax / port_pax) * 100))}%"


def _year_pax_totals(
    qs,
    *,
    years: list[int],
    pax_basis: str,
) -> list[dict[str, Any]]:
    buckets: dict[int, int] = {y: 0 for y in years}
    for booking in qs.iterator(chunk_size=500):
        year = booking.call_date.year if booking.call_date else None
        if year not in buckets:
            continue
        buckets[year] += booking_pax(booking, pax_basis=pax_basis)
    return [{"year": y, "pax": buckets[y]} for y in years]


def build_solicitudes_port_report(
    *,
    date_from: date,
    date_to: date,
    port_id: int,
    years: list[int],
    tag_ids: list[int] | None = None,
    shipping_line_ids: list[int] | None = None,
    without_lta: bool = False,
    pax_basis: str = "planned",
    allowed_ports: set[int] | list[int] | None = None,
    request=None,
) -> dict[str, Any]:
    """
    Left: occupancy rows per selected calendar year (optional tag/line OR filters).
    Right: mismas filas filtradas + totales del puerto (sin naviera) + totales naviera.
    Same status set as other matrix reports (excludes cancelled).
    """
    pax_basis = normalize_pax_basis(pax_basis)
    tag_ids = list(tag_ids or [])
    shipping_line_ids = list(shipping_line_ids or [])
    years = [y for y in years if y >= MIN_REPORT_YEAR]
    if not years:
        years = [
            y
            for y in years_in_range(date_from, date_to)
            if y >= MIN_REPORT_YEAR
        ]

    port = Port.objects.filter(pk=port_id).first()
    if port is None:
        raise ValueError("Puerto no encontrado.")

    summary_years = [
        y for y in years_in_range(date_from, date_to) if y >= MIN_REPORT_YEAR
    ]
    if not summary_years:
        summary_years = list(years)

    tags = list(BookingTag.objects.filter(pk__in=tag_ids).order_by("name"))
    tag_label = ", ".join(t.name for t in tags) if tags else "Tags"

    lines = list(
        ShippingLine.objects.filter(pk__in=shipping_line_ids).order_by("name")
    )
    carrier_label = ", ".join(
        (line.name or line.code or f"#{line.pk}").strip() for line in lines
    )

    base_kwargs = dict(
        date_from=date_from,
        date_to=date_to,
        port_id=port_id,
        allowed_ports=set(allowed_ports) if allowed_ports is not None else None,
        without_lta=without_lta,
    )

    # Port summary: full occupancy set at port (no tag / naviera filter).
    port_qs = scheduled_bookings_qs(**base_kwargs).select_related("vessel")
    port_by_year = _year_pax_totals(port_qs, years=summary_years, pax_basis=pax_basis)

    # Carrier summary: occupancy at port for selected shipping lines only.
    carrier_by_year: list[dict[str, Any]] = []
    if shipping_line_ids:
        carrier_qs = port_qs.filter(shipping_line_id__in=shipping_line_ids)
        carrier_by_year = _year_pax_totals(
            carrier_qs, years=summary_years, pax_basis=pax_basis
        )
        carrier_by_year = [row for row in carrier_by_year if row.get("pax", 0) > 0]

    # Left detail + nuevas: optional tags + shipping lines (OR within each).
    detail_qs = (
        scheduled_bookings_qs(**base_kwargs)
        .select_related("vessel", "port", "shipping_line", "tag")
        .order_by("call_date", "eta", "booking_code")
    )
    if tag_ids:
        detail_qs = detail_qs.filter(tag_id__in=tag_ids)
    if shipping_line_ids:
        detail_qs = detail_qs.filter(shipping_line_id__in=shipping_line_ids)

    year_set = set(years)
    rows_by_year: dict[int, list[dict[str, Any]]] = defaultdict(list)
    nuevas_totals: dict[int, int] = {y: 0 for y in years}

    for booking in detail_qs.iterator(chunk_size=500):
        if not booking.call_date:
            continue
        year = booking.call_date.year
        if year not in year_set:
            continue
        pax = booking_pax(booking, pax_basis=pax_basis)
        rows_by_year[year].append(
            {
                "booking_id": booking.id,
                "booking_code": booking.booking_code,
                "ship": booking.vessel.name if booking.vessel_id else "—",
                "port": _port_label(booking.port) if booking.port_id else _port_label(port),
                "arrival": booking.call_date.isoformat(),
                "arrival_label": _format_arrival(booking.call_date),
                "eta": _format_time(booking.eta),
                "etd": _format_time(booking.etd),
                "pax": pax,
            }
        )
        nuevas_totals[year] += pax

    year_blocks = []
    for year in years:
        rows = rows_by_year.get(year, [])
        if not rows:
            continue
        year_blocks.append(
            {
                "year": year,
                "title": f"Solicitudes {year}",
                "rows": rows,
                "pax_total": nuevas_totals.get(year, 0),
            }
        )

    nuevas = [
        {"year": y, "pax": nuevas_totals.get(y, 0)}
        for y in years
        if nuevas_totals.get(y, 0) > 0
    ]
    nuevas_total = sum(item["pax"] for item in nuevas)

    port_by_year = [row for row in port_by_year if row.get("pax", 0) > 0]

    return {
        "kind": "solicitudes_port",
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "without_lta": without_lta,
        "pax_basis": pax_basis,
        "pax_basis_note": pax_basis_note(pax_basis),
        "port_id": port.id,
        "port_name": _port_label(port),
        "port_logo": _media_url(request, port.logo),
        "years": years,
        "summary_years": summary_years,
        "tag_ids": [t.id for t in tags],
        "tag_names": [t.name for t in tags],
        "tag_label": tag_label,
        "shipping_line_ids": [line.id for line in lines],
        "shipping_line_names": [
            (line.name or line.code or f"#{line.pk}").strip() for line in lines
        ],
        "shipping_line_label": carrier_label,
        "title": _port_label(port).upper(),
        "subtitle": carrier_label or (tag_label if tags else ""),
        "year_blocks": year_blocks,
        "nuevas_solicitadas": nuevas,
        "nuevas_total": nuevas_total,
        "port_by_year": port_by_year,
        "carrier_by_year": carrier_by_year,
    }


def solicitudes_port_filename(port_code: str, date_from: date, date_to: date) -> str:
    _ = (port_code, date_from, date_to)
    return "Resumen de movimientos.xlsx"


def build_solicitudes_port_xlsx(payload: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"

    title = str(payload.get("title") or "RESUMEN")
    subtitle = str(payload.get("subtitle") or "")
    write_title_row(ws, 1, title, 6)
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
        cell = ws.cell(row=2, column=1, value=subtitle)
        style_cell(
            cell,
            font=Font(name="Calibri", size=12, bold=True, color=NAVY),
            fill=FILL_TITLE,
            alignment=ALIGN_LEFT,
            border=None,
        )
        row = 4
    else:
        row = 3

    left_start = row

    # Left year blocks (cols 1-6)
    for block in payload.get("year_blocks") or []:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=1, value=block.get("title") or "")
        style_cell(
            cell,
            font=FONT_TOTAL,
            alignment=ALIGN_LEFT,
            border=None,
        )
        row += 1

        headers = (
            "Ship",
            "Port",
            "Arrival",
            "Hora llegada",
            "Hora salida",
            "Pax",
        )
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=text)
            style_cell(
                cell,
                font=FONT_HEADER,
                fill=FILL_HEADER,
                alignment=ALIGN_CENTER if col > 2 else ALIGN_LEFT,
                border=BORDER_BLACK,
            )
        row += 1

        for item in block.get("rows") or []:
            values = [
                item.get("ship"),
                item.get("port"),
                item.get("arrival_label"),
                item.get("eta"),
                item.get("etd"),
                item.get("pax"),
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                align = ALIGN_LEFT if col <= 2 else ALIGN_CENTER
                if col == 6:
                    align = ALIGN_RIGHT
                style_cell(
                    cell,
                    font=FONT_DATA,
                    fill=FILL_ALT,
                    alignment=align,
                    border=BORDER_BLACK,
                    number_format="#,##0" if col == 6 else None,
                )
            row += 1

        for col in range(1, 6):
            cell = ws.cell(row=row, column=col, value="")
            style_cell(cell, border=None)
        cell = ws.cell(row=row, column=6, value=int(block.get("pax_total") or 0))
        style_cell(
            cell,
            font=FONT_TOTAL,
            fill=FILL_TOTAL,
            alignment=ALIGN_RIGHT,
            border=BORDER_BLACK,
            number_format="#,##0",
        )
        row += 2

    left_end = row

    # Right summary card: port / carrier / Año | Naviera vs total (pct)
    right_row = left_start
    port_name = str(payload.get("port_name") or "Puerto").strip()
    carrier_label = str(payload.get("shipping_line_label") or "").strip()
    carrier_rows = list(
        payload.get("carrier_by_year") or payload.get("nuevas_solicitadas") or []
    )
    port_rows = list(payload.get("port_by_year") or [])
    carrier_by_year = {
        int(item.get("year") or 0): int(item.get("pax") or 0)
        for item in carrier_rows
        if item.get("year")
    }
    port_by_year_map = {
        int(item.get("year") or 0): int(item.get("pax") or 0)
        for item in port_rows
        if item.get("year")
    }
    summary_years = sorted(set(carrier_by_year) | set(port_by_year_map))
    if carrier_label and summary_years:
        for col in (8, 9):
            cell = ws.cell(row=right_row, column=col, value="")
            style_cell(cell, fill=FILL_TITLE, border=BORDER_BLACK)
        cell = ws.cell(row=right_row, column=8, value=port_name)
        style_cell(
            cell,
            font=FONT_TOTAL,
            fill=FILL_TITLE,
            alignment=ALIGN_LEFT,
            border=BORDER_BLACK,
        )
        ws.merge_cells(
            start_row=right_row,
            start_column=8,
            end_row=right_row,
            end_column=9,
        )
        right_row += 1

        for col in (8, 9):
            cell = ws.cell(row=right_row, column=col, value="")
            style_cell(cell, fill=FILL_TITLE, border=BORDER_BLACK)
        cell = ws.cell(row=right_row, column=8, value=carrier_label)
        style_cell(
            cell,
            font=FONT_DATA,
            fill=FILL_TITLE,
            alignment=ALIGN_LEFT,
            border=BORDER_BLACK,
        )
        ws.merge_cells(
            start_row=right_row,
            start_column=8,
            end_row=right_row,
            end_column=9,
        )
        right_row += 1

        cell = ws.cell(row=right_row, column=8, value="Año")
        style_cell(
            cell,
            font=FONT_HEADER,
            fill=FILL_HEADER,
            alignment=ALIGN_CENTER,
            border=BORDER_BLACK,
        )
        cell = ws.cell(
            row=right_row,
            column=9,
            value="Naviera vs total",
        )
        style_cell(
            cell,
            font=FONT_HEADER,
            fill=FILL_HEADER,
            alignment=ALIGN_CENTER,
            border=BORDER_BLACK,
        )
        right_row += 1

        # Help note (UI tooltip equivalent)
        ws.merge_cells(
            start_row=right_row,
            start_column=8,
            end_row=right_row,
            end_column=9,
        )
        cell = ws.cell(
            row=right_row,
            column=8,
            value="Naviera seleccionada vs total de navieras",
        )
        style_cell(
            cell,
            font=FONT_NOTE,
            fill=FILL_ALT,
            alignment=ALIGN_LEFT,
            border=BORDER_BLACK,
        )
        right_row += 1

        carrier_total = 0
        port_total = 0
        for year in summary_years:
            carrier_pax = carrier_by_year.get(year, 0)
            port_pax = port_by_year_map.get(year, 0)
            carrier_total += carrier_pax
            port_total += port_pax
            cell = ws.cell(row=right_row, column=8, value=year)
            style_cell(
                cell,
                font=FONT_DATA,
                fill=FILL_ALT,
                alignment=ALIGN_CENTER,
                border=BORDER_BLACK,
            )
            cell = ws.cell(
                row=right_row,
                column=9,
                value=(
                    f"{carrier_pax:,} / {port_pax:,} "
                    f"({_pct(carrier_pax, port_pax)})"
                ),
            )
            style_cell(
                cell,
                font=FONT_DATA,
                fill=FILL_ALT,
                alignment=ALIGN_RIGHT,
                border=BORDER_BLACK,
            )
            right_row += 1

        cell = ws.cell(row=right_row, column=8, value="Total")
        style_cell(
            cell,
            font=FONT_TOTAL,
            fill=FILL_TOTAL,
            alignment=ALIGN_LEFT,
            border=BORDER_BLACK,
        )
        cell = ws.cell(
            row=right_row,
            column=9,
            value=(
                f"{carrier_total:,} / {port_total:,} "
                f"({_pct(carrier_total, port_total)})"
            ),
        )
        style_cell(
            cell,
            font=FONT_TOTAL,
            fill=FILL_TOTAL,
            alignment=ALIGN_RIGHT,
            border=BORDER_BLACK,
        )

    widths = [22, 16, 12, 12, 12, 10, 3, 14, 28]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    _ = left_end

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
