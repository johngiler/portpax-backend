"""Parse a booking recap sheet (Naviera, Ship, Port, date, ETA/ETD — no position)."""

from __future__ import annotations

from datetime import datetime
from typing import Any, Iterable

from openpyxl import load_workbook

from apps.bookings.services.import_mass.parse_dates import parse_flexible_datetime
from apps.bookings.services.import_mass.parse_itm import (
    ItmParseError,
    _ASSIGNMENT,
    _ARRIVAL,
    _ARRIVAL_DATE_LOOSE,
    _ARRIVAL_DATE_STRICT,
    _DEPARTURE,
    _ETA,
    _ETD,
    _NAVIERA,
    _PORT,
    _SHIP,
    _as_time,
    _cell_str,
    _clock_from_datetime,
    _index,
    header_key,
)


def _parse_table(
    headers: list[str],
    body_rows: Iterable[tuple[int, list[Any]]],
) -> list[dict[str, Any]]:
    header_map = {header_key(h): i for i, h in enumerate(headers) if h}
    # Recap never uses Assignment; ignore if pasted by mistake.
    _ = _index(header_map, _ASSIGNMENT)
    ship_i = _index(header_map, _SHIP)
    port_i = _index(header_map, _PORT)
    naviera_i = _index(header_map, _NAVIERA)
    date_strict_i = _index(header_map, _ARRIVAL_DATE_STRICT)
    date_loose_i = _index(header_map, _ARRIVAL_DATE_LOOSE)
    eta_i = _index(header_map, _ETA)
    etd_i = _index(header_map, _ETD)
    arrival_i = _index(header_map, _ARRIVAL)
    departure_i = _index(header_map, _DEPARTURE)

    date_i = date_strict_i
    if date_i is None and arrival_i is None:
        date_i = date_loose_i

    if ship_i is None or port_i is None or (date_i is None and arrival_i is None):
        raise ItmParseError(
            "Faltan columnas. Usa Naviera, Ship, Port, Arrival Date, ETA y ETD "
            "(sin posición). También vale Ship, Port, Arrival y Departure."
        )

    parsed: list[dict[str, Any]] = []
    for excel_row, values in body_rows:
        def cell(index: int | None) -> Any:
            if index is None or index >= len(values):
                return None
            return values[index]

        ship = _cell_str(cell(ship_i))
        port = _cell_str(cell(port_i))
        if not ship and not port:
            continue

        line_raw = _cell_str(cell(naviera_i)) if naviera_i is not None else ""
        berth = parse_flexible_datetime(cell(date_i)) if date_i is not None else None
        arrival = (
            parse_flexible_datetime(cell(arrival_i)) if arrival_i is not None else None
        )
        departure = (
            parse_flexible_datetime(cell(departure_i))
            if departure_i is not None
            else None
        )
        call_dt = berth or arrival
        if call_dt is None:
            parsed.append(
                {
                    "row_number": excel_row,
                    "ship": ship,
                    "port_raw": port,
                    "line_raw": line_raw,
                    "call_date": None,
                    "eta": None,
                    "etd": None,
                }
            )
            continue

        eta = _as_time(cell(eta_i)) if eta_i is not None else None
        etd = _as_time(cell(etd_i)) if etd_i is not None else None
        if eta is None:
            eta = _clock_from_datetime(arrival, explicit=False)
        if etd is None:
            etd = _clock_from_datetime(departure, explicit=False)

        parsed.append(
            {
                "row_number": excel_row,
                "ship": ship,
                "port_raw": port,
                "line_raw": line_raw,
                "call_date": call_dt.date() if isinstance(call_dt, datetime) else call_dt,
                "eta": eta,
                "etd": etd,
            }
        )

    if not parsed:
        raise ItmParseError("No se encontraron filas de recap.")
    return parsed


def parse_recap_workbook(file_obj) -> list[dict[str, Any]]:
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration as exc:
            raise ItmParseError("El archivo está vacío.") from exc
        headers = [_cell_str(h) for h in (header or ())]
        body = [
            (excel_row, list(row or ()))
            for excel_row, row in enumerate(rows_iter, start=2)
            if row is not None
        ]
        return _parse_table(headers, body)
    finally:
        wb.close()


def parse_recap_tsv(text: str) -> list[dict[str, Any]]:
    raw = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not raw:
        raise ItmParseError(
            "Pega al menos una fila con Naviera, Ship, Port y Arrival Date."
        )
    lines = [ln for ln in raw.split("\n") if ln.strip()]

    def split_line(line: str) -> list[str]:
        if "\t" in line:
            return [_cell_str(c) for c in line.split("\t")]
        if ";" in line:
            return [_cell_str(c) for c in line.split(";")]
        return [_cell_str(line)]

    headers = split_line(lines[0])
    body = [(i, split_line(line)) for i, line in enumerate(lines[1:], start=2)]
    if not body:
        raise ItmParseError(
            "Incluye la fila de encabezados y al menos una fila de datos."
        )
    return _parse_table(headers, body)
